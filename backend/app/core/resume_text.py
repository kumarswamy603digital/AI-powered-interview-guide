from __future__ import annotations

"""
Resume text extraction.

Uploads previously landed on disk and were never read, which meant every
downstream feature (ATS scoring, interviews, ranking) required a human to paste
resume text by hand. This module turns an uploaded file into text at upload time.

Optional dependencies follow the same pattern as the Gemini integration: if the
parser library is missing the upload still succeeds, but the extraction status
records why no text is available instead of failing silently.
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

try:  # Optional dependency for PDF parsing
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional
    PdfReader = None  # type: ignore[assignment]

try:  # Optional dependency for DOCX parsing
    import docx  # python-docx
except Exception:  # pragma: no cover - optional
    docx = None  # type: ignore[assignment]


PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_CONTENT_TYPE = "text/plain"


# Extraction status values persisted on the resume row.
STATUS_OK = "ok"
STATUS_EMPTY = "empty"
STATUS_UNSUPPORTED = "unsupported"
STATUS_MISSING_DEPENDENCY = "missing_dependency"
STATUS_FAILED = "failed"


@dataclass
class ExtractionResult:
    text: str
    status: str
    detail: Optional[str] = None

    @property
    def ok(self) -> bool:
        return self.status == STATUS_OK


def _normalize_whitespace(text: str) -> str:
    lines = [line.strip() for line in (text or "").splitlines()]
    # Drop runs of blank lines so the text stays compact in LLM prompts.
    cleaned: list[str] = []
    for line in lines:
        if not line and cleaned and not cleaned[-1]:
            continue
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def _extract_pdf(path: Path) -> ExtractionResult:
    if PdfReader is None:
        return ExtractionResult(
            text="",
            status=STATUS_MISSING_DEPENDENCY,
            detail="pypdf is not installed; run pip install -r requirements.txt",
        )
    try:
        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = _normalize_whitespace("\n".join(pages))
    except Exception as exc:  # pragma: no cover - corrupt/encrypted files
        return ExtractionResult(text="", status=STATUS_FAILED, detail=str(exc)[:300])

    if not text:
        return ExtractionResult(
            text="",
            status=STATUS_EMPTY,
            detail="No selectable text found (the PDF may be a scan).",
        )
    return ExtractionResult(text=text, status=STATUS_OK)


def _extract_docx(path: Path) -> ExtractionResult:
    if docx is None:
        return ExtractionResult(
            text="",
            status=STATUS_MISSING_DEPENDENCY,
            detail="python-docx is not installed; run pip install -r requirements.txt",
        )
    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        # Skills are very often laid out in tables, so read those too.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = _normalize_whitespace("\n".join(parts))
    except Exception as exc:  # pragma: no cover - corrupt files
        return ExtractionResult(text="", status=STATUS_FAILED, detail=str(exc)[:300])

    if not text:
        return ExtractionResult(text="", status=STATUS_EMPTY, detail="Document contained no text.")
    return ExtractionResult(text=text, status=STATUS_OK)


def _extract_plain(path: Path) -> ExtractionResult:
    try:
        raw = path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:  # pragma: no cover
        return ExtractionResult(text="", status=STATUS_FAILED, detail=str(exc)[:300])
    text = _normalize_whitespace(raw)
    if not text:
        return ExtractionResult(text="", status=STATUS_EMPTY, detail="File was empty.")
    return ExtractionResult(text=text, status=STATUS_OK)


def extract_resume_text(storage_path: str | Path, content_type: str) -> ExtractionResult:
    """
    Extract plain text from an uploaded resume.

    Never raises: callers persist the returned status so HR can see which resumes
    need a manual paste instead of hitting an empty ranking with no explanation.
    """
    path = Path(storage_path)
    if not path.exists():
        return ExtractionResult(text="", status=STATUS_FAILED, detail="Stored file not found.")

    suffix = path.suffix.lower()
    if content_type == PDF_CONTENT_TYPE or suffix == ".pdf":
        return _extract_pdf(path)
    if content_type == DOCX_CONTENT_TYPE or suffix == ".docx":
        return _extract_docx(path)
    if content_type.startswith("text/") or suffix in {".txt", ".md"}:
        return _extract_plain(path)

    return ExtractionResult(
        text="",
        status=STATUS_UNSUPPORTED,
        detail=f"Unsupported content type: {content_type}",
    )
